from fastapi import FastAPI, UploadFile, File, Form, HTTPException
import openpyxl
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from pydantic import BaseModel
from fastapi.responses import FileResponse
from sqlalchemy import func
from backend.document_reader import extract_text_from_file
from backend.plagiarism_engine import analyze_submissions
from backend.database import create_tables, SessionLocal
from backend import models
from io import BytesIO
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import os
import shutil
from datetime import datetime, timezone
import json
import uuid

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle
)

# =========================================================
# ROSTER REQUEST MODELS
# =========================================================

class SectionCreate(BaseModel):
    name: str
    faculty_name: str | None = None


class StudentCreate(BaseModel):
    student_id: str
    name: str
    email: str | None = None
    department: str | None = None
    year: str | None = None
    section_id: int


class StudentProfileUpdate(BaseModel):
    name: str
    email: str | None = None
    department: str | None = None
    year: str | None = None


class AssignmentCreate(BaseModel):
    section_id: int
    title: str
    subject: str
    description: str | None = None
    due_date: datetime | None = None
    faculty_name: str | None = None


def serialize_due_date(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")

# ==============================
# Request Models
# ==============================

class AnalysisRequest(BaseModel):
    submission_1: str
    submission_2: str


# ==============================
# FastAPI App
# ==============================

app = FastAPI(
    title="VeriText API",
    description="Plagiarism Detection Engine for Academic Submissions",
    version="1.0.0"
)


# ==============================
# Database
# ==============================

create_tables()


# ==============================
# CORS
# ==============================

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["*"],
)


# ==============================
# Home
# ==============================

@app.get("/")
def home():
    return {
        "message": "VeriText backend is running successfully!"
    }


# ==============================
# Health Check
# ==============================

@app.get("/api/health")
def health_check():
    return {
        "status": "healthy",
        "service": "VeriText API"
    }


# ==============================
# Plagiarism Analysis
# ==============================

@app.post("/api/analyze")
def analyze(request: AnalysisRequest):

    result = analyze_submissions(
        request.submission_1,
        request.submission_2
    )

    return {
        "success": True,
        "similarity_percentage": result["similarity_percentage"],
        "fingerprints_1": result["fingerprints_1"],
        "fingerprints_2": result["fingerprints_2"],
        "winnowed_1": result["winnowed_1"],
        "winnowed_2": result["winnowed_2"],
        "common_fingerprints": result["common_fingerprints"],
        "common_passages": result["common_passages"]
    }


# ==============================
# Create Submission
# ==============================

@app.post("/api/submissions")
async def create_submission(
    student_id: str = Form(...),
    student_name: str = Form(...),
    assignment_title: str | None = Form(None),
    subject: str | None = Form(None),
    assignment_id: int | None = Form(None),
    file: UploadFile = File(...)
):

    db = SessionLocal()

    file_path = None

    try:

        # -------------------------------------------------
        # Find student in roster
        # -------------------------------------------------

        student = (
            db.query(models.Student)
            .filter(
                models.Student.student_id == student_id
            )
            .first()
        )

        if not student:

            raise HTTPException(
                status_code=404,
                detail=(
                    f"Student ID '{student_id}' "
                    "is not present in the roster."
                )
            )

        # -------------------------------------------------
        # Make sure submitted name matches roster
        # -------------------------------------------------

        if (
            student_name.strip().lower()
            != student.name.strip().lower()
        ):

            raise HTTPException(
                status_code=400,
                detail={
                    "message":
                        "Student name does not match "
                        "the roster.",
                    "roster_name":
                        student.name,
                    "provided_name":
                        student_name
                }
            )

        assignment = None
        if assignment_id is not None:
            assignment = (
                db.query(models.Assignment)
                .filter(models.Assignment.id == assignment_id)
                .first()
            )
            if not assignment:
                raise HTTPException(
                    status_code=404,
                    detail="Assignment not found."
                )
            if assignment.section_id != student.section_id:
                raise HTTPException(
                    status_code=403,
                    detail="This assignment is not assigned to your section."
                )
            assignment_title = assignment.title
            subject = assignment.subject

        if not assignment_title or not subject:
            raise HTTPException(
                status_code=400,
                detail="Select a valid assignment before submitting."
            )

        # -------------------------------------------------
        # Check whether this student already submitted
        # this assignment
        # -------------------------------------------------

        existing_submission = (
            db.query(models.Submission)
            .filter(
                models.Submission.student_record_id
                == student.id,
                (
                    models.Submission.assignment_id == assignment.id
                    if assignment
                    else models.Submission.assignment_title == assignment_title
                )
            )
            .first()
        )

        if existing_submission:

            raise HTTPException(
                status_code=400,
                detail=(
                    "This student has already submitted "
                    "this assignment."
                )
            )

        if not file.filename:
            raise HTTPException(status_code=400, detail="No file was selected.")
        safe_filename = os.path.basename(file.filename)
        extension = os.path.splitext(safe_filename)[1].lower()
        if extension not in {".pdf", ".docx", ".txt"}:
            raise HTTPException(
                status_code=400,
                detail="Only PDF, DOCX, and TXT files are supported."
            )
        contents = await file.read()
        if len(contents) > 10 * 1024 * 1024:
            raise HTTPException(
                status_code=413,
                detail="File exceeds the 10 MB upload limit."
            )
        await file.seek(0)

        # -------------------------------------------------
        # Store uploaded files inside backend/uploads
        # -------------------------------------------------

        upload_directory = os.path.join(
            os.path.dirname(__file__),
            "uploads"
        )

        os.makedirs(
            upload_directory,
            exist_ok=True
        )

        # -------------------------------------------------
        # Create file path
        # -------------------------------------------------

        file_path = os.path.join(
            upload_directory,
            f"{uuid.uuid4().hex}_{safe_filename}"
        )

        # -------------------------------------------------
        # Save uploaded file
        # -------------------------------------------------

        with open(file_path, "wb") as buffer:

            shutil.copyfileobj(
                file.file,
                buffer
            )

        # -------------------------------------------------
        # Create submission linked to Student
        # -------------------------------------------------

        submission = models.Submission(

            student_record_id=student.id,

            assignment_id=assignment.id if assignment else None,

            student_name=student.name,

            assignment_title=assignment_title,

            subject=subject,

            filename=safe_filename,

            file_path=file_path,

            status="PENDING",

            submitted_at=datetime.utcnow()
        )

        db.add(submission)

        db.commit()

        db.refresh(submission)

        return {
            "success": True,

            "message":
                "Assignment submitted successfully.",

            "submission_id":
                submission.id,

            "student_id":
                student.student_id,

            "student_name":
                student.name,

            "assignment_title":
                submission.assignment_title,

            "filename":
                submission.filename,

            "status":
                submission.status
        }

    except HTTPException:

        db.rollback()

        # Delete uploaded file if validation failed
        if file_path and os.path.exists(file_path):
            os.remove(file_path)

        raise

    except Exception as e:

        db.rollback()

        # Delete uploaded file if database operation failed
        if file_path and os.path.exists(file_path):
            os.remove(file_path)

        raise HTTPException(
            status_code=500,
            detail=(
                "Assignment submission failed. "
                f"Error: {str(e)}"
            )
        )

    finally:

        db.close()

# ==============================
# Get All Submissions
# ==============================

@app.get("/api/submissions")
def get_submissions():

    db = SessionLocal()

    try:

        submissions = (
            db.query(models.Submission)
            .order_by(
                models.Submission.submitted_at.desc()
            )
            .all()
        )

        return {
            "success": True,
            "submissions": [
                {
                    "id": submission.id,
                    "student_name": submission.student_name,
                    "assignment_title": submission.assignment_title,
                    "subject": submission.subject,
                    "filename": submission.filename,
                    "status": submission.status,
                    "submitted_at": (
                        submission.submitted_at.isoformat()
                        if submission.submitted_at
                        else None
                    )
                }
                for submission in submissions
            ]
        }

    finally:

        db.close()


# ==============================
# Get Single Submission
# ==============================

@app.get("/api/submissions/{submission_id}")
def get_submission(submission_id: int):

    db = SessionLocal()

    try:

        submission = (
            db.query(models.Submission)
            .filter(models.Submission.id == submission_id)
            .first()
        )

        if not submission:
            return {
                "success": False,
                "message": "Submission not found."
            }

        return {
            "success": True,
            "submission": {
                "id": submission.id,
                "student_name": submission.student_name,
                "assignment_title": submission.assignment_title,
                "subject": submission.subject,
                "filename": submission.filename,
                "file_path": submission.file_path,
                "status": submission.status,
                "submitted_at": (
                    submission.submitted_at.isoformat()
                    if submission.submitted_at
                    else None
                )
            }
        }

    finally:

        db.close()

@app.get("/api/submissions/{submission_id}/file")
def get_submission_file(submission_id: int):
    db = SessionLocal()

    try:
        submission = (
            db.query(models.Submission)
            .filter(models.Submission.id == submission_id)
            .first()
        )

        if not submission:
            return {
                "success": False,
                "message": "Submission not found."
            }

        if not submission.file_path:
            return {
                "success": False,
                "message": "File path not available."
            }

        if not os.path.exists(submission.file_path):
            return {
                "success": False,
                "message": "Submitted file not found on server."
            }

        return FileResponse(
            path=submission.file_path,
            filename=submission.filename
        )

    finally:
        db.close()

# ==============================
# Publish Faculty Evaluation
# ==============================

class EvaluationRequest(BaseModel):
    submission_id: int
    marks: float
    similarity_percentage: float
    common_passages: str
    feedback: str
    faculty_name: str

    fingerprints_1: int = 0
    fingerprints_2: int = 0
    winnowed_1: int = 0
    winnowed_2: int = 0
    common_fingerprints: int = 0


@app.post("/api/evaluations")
def create_evaluation(request: EvaluationRequest):

    db = SessionLocal()

    try:

        submission = (
            db.query(models.Submission)
            .filter(
                models.Submission.id == request.submission_id
            )
            .first()
        )

        if not submission:
            return {
                "success": False,
                "message": "Submission not found."
            }

        evaluation = models.Evaluation(
    submission_id=request.submission_id,

    marks=request.marks,

    similarity_percentage=request.similarity_percentage,

    fingerprints_1=request.fingerprints_1,
    fingerprints_2=request.fingerprints_2,

    winnowed_1=request.winnowed_1,
    winnowed_2=request.winnowed_2,

    common_fingerprints=request.common_fingerprints,

    common_passages=request.common_passages,

    feedback=request.feedback,

    faculty_name=request.faculty_name,

    evaluation_status="COMPLETED",

    evaluated_at=datetime.utcnow()
)

        db.add(evaluation)

        submission.status = "EVALUATED"

        db.commit()

        db.refresh(evaluation)

        return {
            "success": True,
            "message": "Evaluation published successfully.",
            "evaluation_id": evaluation.id,
            "submission_id": submission.id,
            "status": submission.status
        }

    finally:

        db.close()

@app.get("/api/evaluations/{submission_id}")
def get_evaluation(submission_id: int):

    db = SessionLocal()

    try:

        # =========================================
        # GET LATEST FACULTY EVALUATION
        # =========================================

        evaluation = (
            db.query(models.Evaluation)
            .filter(
                models.Evaluation.submission_id == submission_id
            )
            .order_by(
                models.Evaluation.evaluated_at.desc()
            )
            .first()
        )

        if not evaluation:

            return {
                "success": False,
                "message": "Evaluation not found."
            }


        # =========================================
        # GET SELECTED SUBMISSION
        # =========================================

        submission = (
            db.query(models.Submission)
            .filter(
                models.Submission.id == submission_id
            )
            .first()
        )

        if not submission:

            return {
                "success": False,
                "message": "Submission not found."
            }


        # =========================================
        # ENGINE ANALYSIS
        # =========================================

        engine_result = None

        if (
            submission.file_path
            and os.path.exists(submission.file_path)
        ):

            try:

                with open(
                    submission.file_path,
                    "r",
                    encoding="utf-8",
                    errors="ignore"
                ) as file:

                    submission_text = file.read()


                # Get other submissions
                other_query = db.query(models.Submission).filter(
                    models.Submission.id != submission_id
                )
                if submission.assignment_id is not None:
                    other_query = other_query.filter(
                        models.Submission.assignment_id == submission.assignment_id
                    )
                else:
                    other_query = other_query.filter(
                        models.Submission.assignment_title == submission.assignment_title,
                        models.Submission.subject == submission.subject
                    )
                other_submissions = other_query.all()


                best_match = None


                for other in other_submissions:

                    if not other.file_path:
                        continue

                    if not os.path.exists(
                        other.file_path
                    ):
                        continue


                    try:

                        with open(
                            other.file_path,
                            "r",
                            encoding="utf-8",
                            errors="ignore"
                        ) as file:

                            other_text = file.read()


                        analysis = analyze_submissions(
                            submission_text,
                            other_text
                        )


                        if (
                            best_match is None
                            or analysis[
                                "similarity_percentage"
                            ]
                            >
                            best_match[
                                "similarity_percentage"
                            ]
                        ):

                            best_match = analysis


                    except Exception:

                        continue


                if best_match:

                    engine_result = best_match


            except Exception as error:

                print(
                    "Engine analysis error:",
                    error
                )


        # =========================================
        # COMMON PASSAGES
        # =========================================

        stored_passages = []

        try:

            if evaluation.common_passages:

                stored_passages = json.loads(
                    evaluation.common_passages
                )

        except Exception as error:

            print(
                "Stored passage parsing error:",
                error
            )

            stored_passages = []


        # Prefer stored faculty-evaluation passages.
        # If none exist, use fresh engine passages.

        if not stored_passages and engine_result:

            stored_passages = (
                engine_result.get(
                    "common_passages",
                    []
                )
            )
                    # =========================================
        # ENGINE METRICS
        # =========================================

        fingerprints_1 = evaluation.fingerprints_1
        fingerprints_2 = evaluation.fingerprints_2

        winnowed_1 = evaluation.winnowed_1
        winnowed_2 = evaluation.winnowed_2

        common_fingerprints = evaluation.common_fingerprints


        # =========================================
        # RETURN COMPLETE EVALUATION
        # =========================================

        return {

            "success": True,

            "evaluation": {

                "id":
                    evaluation.id,

                "submission_id":
                    evaluation.submission_id,

                "marks":
                    evaluation.marks,

                "similarity_percentage":
                    evaluation.similarity_percentage,

            "fingerprints_1":
    fingerprints_1,

"fingerprints_2":
    fingerprints_2,

"winnowed_1":
    winnowed_1,

"winnowed_2":
    winnowed_2,

"common_fingerprints":
    common_fingerprints,
                "common_matches":
                    len(stored_passages),

                "passage_count":
                    len(stored_passages),

                "common_passages":
                    json.dumps(stored_passages),

                "feedback":
                    evaluation.feedback,

                "faculty_name":
                    evaluation.faculty_name,

                "evaluation_status":
                    evaluation.evaluation_status,

                "evaluated_at":
                    (
                        evaluation.evaluated_at.isoformat()
                        if evaluation.evaluated_at
                        else None
                    )
            }
        }


    finally:

        db.close()
@app.get("/api/evaluations")
def get_all_evaluations():
    db = SessionLocal()

    try:
        evaluations = (
            db.query(models.Evaluation)
            .order_by(models.Evaluation.evaluated_at.desc())
            .all()
        )

        results = []

        for evaluation in evaluations:

            submission = (
                db.query(models.Submission)
                .filter(
                    models.Submission.id ==
                    evaluation.submission_id
                )
                .first()
            )

            if not submission:
                continue

            results.append({
                "evaluation_id": evaluation.id,
                "submission_id": submission.id,
                "student_name": submission.student_name,
                "assignment_title": submission.assignment_title,
                "subject": submission.subject,
                "filename": submission.filename,
                "marks": evaluation.marks,
                "similarity_percentage":
                    evaluation.similarity_percentage,
                "evaluation_status":
                    evaluation.evaluation_status,
                "faculty_name":
                    evaluation.faculty_name,
                "evaluated_at":
                    evaluation.evaluated_at.isoformat()
                    if evaluation.evaluated_at
                    else None
            })

        return {
            "success": True,
            "evaluations": results
        }

    finally:
        db.close()

@app.get("/api/reports/{submission_id}/pdf")
def download_report_pdf(submission_id: int):

    db = SessionLocal()

    try:
        submission = (
            db.query(models.Submission)
            .filter(models.Submission.id == submission_id)
            .first()
        )

        if not submission:
            return {
                "success": False,
                "message": "Submission not found."
            }

        evaluation = (
            db.query(models.Evaluation)
            .filter(models.Evaluation.submission_id == submission_id)
            .order_by(models.Evaluation.evaluated_at.desc())
            .first()
        )

        if not evaluation:
            return {
                "success": False,
                "message": "Evaluation not found."
            }

        reports_dir = os.path.join(
            os.path.dirname(__file__),
            "reports"
        )

        os.makedirs(reports_dir, exist_ok=True)

        pdf_path = os.path.join(
            reports_dir,
            f"VeriText_Report_{submission_id}.pdf"
        )

        doc = SimpleDocTemplate(
            pdf_path,
            pagesize=A4,
            rightMargin=45,
            leftMargin=45,
            topMargin=45,
            bottomMargin=45
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            fontSize=22,
            spaceAfter=8
        )

        subtitle_style = ParagraphStyle(
            "Subtitle",
            parent=styles["Normal"],
            alignment=TA_CENTER,
            fontSize=9,
            textColor=colors.grey,
            spaceAfter=25
        )

        heading_style = ParagraphStyle(
            "Heading",
            parent=styles["Heading2"],
            fontSize=12,
            spaceBefore=16,
            spaceAfter=8
        )

        body_style = ParagraphStyle(
            "Body",
            parent=styles["BodyText"],
            fontSize=9,
            leading=14
        )

        story = []

        story.append(
            Paragraph(
                "VERITEXT",
                title_style
            )
        )

        story.append(
            Paragraph(
                "PLAGIARISM & EVALUATION REPORT",
                subtitle_style
            )
        )

        assignment_data = [
            ["Student", str(submission.student_name)],
            ["Assignment", str(submission.assignment_title)],
            ["Subject", str(submission.subject)],
            ["Submission ID", f"S-{submission.id:04d}"],
            ["File", str(submission.filename)],
        ]

        assignment_table = Table(
            assignment_data,
            colWidths=[110, 360]
        )

        assignment_table.setStyle(
            TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ])
        )

        story.append(assignment_table)

        story.append(
            Paragraph(
                "Evaluation Summary",
                heading_style
            )
        )

        summary_data = [
            ["Marks", f"{evaluation.marks or 0} / 100"],
            [
                "Similarity",
                f"{evaluation.similarity_percentage or 0}%"
            ],
            [
                "Status",
                str(evaluation.evaluation_status or "COMPLETED")
            ],
            [
                "Faculty",
                str(evaluation.faculty_name or "Faculty")
            ],
        ]

        summary_table = Table(
            summary_data,
            colWidths=[110, 360]
        )

        summary_table.setStyle(
            TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ])
        )

        story.append(summary_table)

        story.append(
            Paragraph(
                "Faculty Feedback",
                heading_style
            )
        )

        feedback = evaluation.feedback or "No feedback provided."

        story.append(
            Paragraph(
                feedback,
                body_style
            )
        )

        story.append(
            Paragraph(
                "Matching Content",
                heading_style
            )
        )

        passages = []

        try:
            if evaluation.common_passages:
                passages = json.loads(
                    evaluation.common_passages
                )
        except Exception:
            passages = []

        if passages:

            for index, passage in enumerate(passages):

                passage_1 = passage.get(
                    "submission_1_passage",
                    ""
                )

                passage_2 = passage.get(
                    "submission_2_passage",
                    ""
                )

                length = passage.get(
                    "length",
                    ""
                )

                story.append(
                    Paragraph(
                        f"<b>Match {index + 1:02d} — "
                        f"{length} words</b>",
                        body_style
                    )
                )

                match_data = [
                    [
                        Paragraph(
                            "<b>Submission</b><br/>"
                            + passage_1,
                            body_style
                        ),
                        Paragraph(
                            "<b>Matched Content</b><br/>"
                            + passage_2,
                            body_style
                        )
                    ]
                ]

                match_table = Table(
                    match_data,
                    colWidths=[230, 230]
                )

                match_table.setStyle(
                    TableStyle([
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BACKGROUND", (0, 0), (-1, -1), colors.whitesmoke),
                        ("LEFTPADDING", (0, 0), (-1, -1), 8),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                        ("TOPPADDING", (0, 0), (-1, -1), 8),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                    ])
                )

                story.append(match_table)
                story.append(Spacer(1, 12))

        else:

            story.append(
                Paragraph(
                    "No significant matching passages were recorded.",
                    body_style
                )
            )

        story.append(
            Spacer(1, 20)
        )

        story.append(
            Paragraph(
                "Generated by VeriText — Plagiarism Detection Engine",
                subtitle_style
            )
        )

        doc.build(story)

        return FileResponse(
            pdf_path,
            media_type="application/pdf",
            filename=f"VeriText_Report_{submission_id}.pdf"
        )

    finally:
        db.close()

@app.get("/api/submissions/{submission_id}/document")
def get_submission_document(submission_id: int):

    db = SessionLocal()

    try:

        submission = (
            db.query(models.Submission)
            .filter(
                models.Submission.id == submission_id
            )
            .first()
        )

        if not submission:

            return {
                "success": False,
                "message": "Submission not found."
            }


        if not submission.file_path:

            return {
                "success": False,
                "message": "File path not available."
            }


        try:

            text = extract_text_from_file(
                submission.file_path
            )

        except Exception as error:

            return {
                "success": False,
                "message": f"Unable to read document: {error}"
            }


        return {

            "success": True,

            "document": {

                "id": submission.id,

                "student_name":
                    submission.student_name,

                "assignment_title":
                    submission.assignment_title,

                "subject":
                    submission.subject,

                "filename":
                    submission.filename,

                "text":
                    text

            }

        }

    finally:

        db.close()
# ==============================
# Serve Submitted Assignment
# ==============================

@app.get("/api/submissions/{submission_id}/content")
def get_submission_content(submission_id: int):
    db = SessionLocal()

    try:
        submission = (
            db.query(models.Submission)
            .filter(models.Submission.id == submission_id)
            .first()
        )

        if not submission:
            return {
                "success": False,
                "message": "Submission not found."
            }

        if not submission.file_path:
            return {
                "success": False,
                "message": "Submission file path not available."
            }

        if not os.path.exists(submission.file_path):
            return {
                "success": False,
                "message": "Submitted file not found on server."
            }

        with open(
            submission.file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:
            content = file.read()

        return {
            "success": True,
            "content": content
        }

    finally:
        db.close()

@app.post("/api/submissions/{submission_id}/analyze")
def analyze_submission(submission_id: int):

    db = SessionLocal()

    try:
        # Get the selected submission
        submission = (
            db.query(models.Submission)
            .filter(models.Submission.id == submission_id)
            .first()
        )

        if not submission:
            return {
                "success": False,
                "message": "Submission not found."
            }

        # Make sure the selected file exists
        if not submission.file_path:
            return {
                "success": False,
                "message": "Submission file path not available."
            }

        if not os.path.exists(submission.file_path):
            return {
                "success": False,
                "message": "Submission file not found."
            }

        # Read selected submission
        with open(
            submission.file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:
            submission_text = file.read()

        # Get all other submissions
        other_query = db.query(models.Submission).filter(
            models.Submission.id != submission_id
        )
        if submission.assignment_id is not None:
            other_query = other_query.filter(
                models.Submission.assignment_id == submission.assignment_id
            )
        else:
            other_query = other_query.filter(
                models.Submission.assignment_title == submission.assignment_title,
                models.Submission.subject == submission.subject
            )
        other_submissions = other_query.all()

        results = []

        for other in other_submissions:

            if not other.file_path:
                continue

            if not os.path.exists(other.file_path):
                continue

            # Read other submission
            with open(
                other.file_path,
                "r",
                encoding="utf-8",
                errors="ignore"
            ) as file:
                other_text = file.read()

            # Run existing plagiarism engine
            analysis = analyze_submissions(
                submission_text,
                other_text
            )

            results.append({
                "submission_id": other.id,
                "student_name": other.student_name,
                "assignment_title": other.assignment_title,
                "similarity_percentage":
                    analysis["similarity_percentage"],
                "fingerprints_1":
                    analysis["fingerprints_1"],
                "fingerprints_2":
                    analysis["fingerprints_2"],
                "winnowed_1":
                    analysis["winnowed_1"],
                "winnowed_2":
                    analysis["winnowed_2"],
                "common_fingerprints":
                    analysis["common_fingerprints"],
                "common_passages":
                    analysis["common_passages"]
            })

        # Sort highest similarity first
        results.sort(
            key=lambda item: item["similarity_percentage"],
            reverse=True
        )

        # No comparison submissions available
        if not results:
            return {
                "success": True,
                "message": "No other submissions available for comparison.",
                "submission": {
                    "id": submission.id,
                    "student_name": submission.student_name,
                    "assignment_title": submission.assignment_title
                },
                "results": []
            }

        # Highest similarity result
        best_match = results[0]

        return {
            "success": True,

            "submission": {
                "id": submission.id,
                "student_name": submission.student_name,
                "assignment_title": submission.assignment_title,
                "subject": submission.subject,
                "filename": submission.filename
            },

            "similarity_percentage":
                best_match["similarity_percentage"],

            "matched_submission": {
                "id": best_match["submission_id"],
                "student_name": best_match["student_name"],
                "assignment_title":
                    best_match["assignment_title"]
            },

            "fingerprints_1":
                best_match["fingerprints_1"],

            "fingerprints_2":
                best_match["fingerprints_2"],

            "winnowed_1":
                best_match["winnowed_1"],

            "winnowed_2":
                best_match["winnowed_2"],

            "common_fingerprints":
                best_match["common_fingerprints"],

            "common_passages":
                best_match["common_passages"],

            "all_matches": results
        }

    finally:
        db.close()
        # =========================================================
# SECTION APIs
# =========================================================

@app.post("/api/sections")
def create_section(section: SectionCreate):

    db = SessionLocal()

    try:
        section_name = section.name.strip()
        if not section_name:
            raise HTTPException(status_code=400, detail="Section name is required.")

        # Prevent duplicate section names
        existing_section = (
            db.query(models.Section)
            .filter(
                func.lower(models.Section.name) == section_name.lower()
            )
            .first()
        )

        if existing_section:
            raise HTTPException(
                status_code=400,
                detail="Section already exists."
            )

        new_section = models.Section(
            name=section_name,
            faculty_name=section.faculty_name.strip() if section.faculty_name else None
        )

        db.add(new_section)
        db.commit()
        db.refresh(new_section)

        return {
            "success": True,
            "message": "Section created successfully.",
            "section": {
                "id": new_section.id,
                "name": new_section.name,
                "faculty_name": new_section.faculty_name
            }
        }

    finally:
        db.close()


@app.put("/api/sections/{section_id}")
def update_section(section_id: int, section: SectionCreate):
    db = SessionLocal()
    try:
        existing_section = (
            db.query(models.Section)
            .filter(models.Section.id == section_id)
            .first()
        )
        if not existing_section:
            raise HTTPException(status_code=404, detail="Section not found.")
        section_name = section.name.strip()
        if not section_name:
            raise HTTPException(status_code=400, detail="Section name is required.")
        duplicate = (
            db.query(models.Section)
            .filter(
                func.lower(models.Section.name) == section_name.lower(),
                models.Section.id != section_id
            )
            .first()
        )
        if duplicate:
            raise HTTPException(status_code=409, detail="Section already exists.")
        existing_section.name = section_name
        existing_section.faculty_name = (
            section.faculty_name.strip() if section.faculty_name else None
        )
        db.commit()
        db.refresh(existing_section)
        return {
            "success": True,
            "section": {
                "id": existing_section.id,
                "name": existing_section.name,
                "faculty_name": existing_section.faculty_name,
                "student_count": len(existing_section.students)
            }
        }
    finally:
        db.close()


@app.get("/api/sections")
def get_sections():

    db = SessionLocal()

    try:

        sections = (
            db.query(models.Section)
            .order_by(models.Section.id)
            .all()
        )

        return {
            "success": True,
            "count": len(sections),
            "sections": [
                {
                    "id": section.id,
                    "name": section.name,
                    "faculty_name": section.faculty_name,
                    "student_count": len(section.students)
                }
                for section in sections
            ]
        }

    finally:
        db.close()


@app.get("/api/sections/{section_id}")
def get_section(section_id: int):

    db = SessionLocal()

    try:

        section = (
            db.query(models.Section)
            .filter(
                models.Section.id == section_id
            )
            .first()
        )

        if not section:
            raise HTTPException(
                status_code=404,
                detail="Section not found."
            )

        return {
            "success": True,
            "section": {
                "id": section.id,
                "name": section.name,
                "faculty_name": section.faculty_name,
                "student_count": len(section.students)
            }
        }

    finally:
        db.close()
        # =========================================================
# STUDENT APIs
# =========================================================

@app.post("/api/students")
def create_student(student: StudentCreate):

    db = SessionLocal()

    try:

        # Check section
        section = (
            db.query(models.Section)
            .filter(
                models.Section.id == student.section_id
            )
            .first()
        )

        if not section:
            raise HTTPException(
                status_code=404,
                detail="Section not found."
            )

        # Prevent duplicate student ID
        existing_student = (
            db.query(models.Student)
            .filter(
                models.Student.student_id == student.student_id
            )
            .first()
        )

        if existing_student:
            raise HTTPException(
                status_code=400,
                detail="Student ID already exists."
            )

        if not student.student_id.strip() or not student.name.strip():
            raise HTTPException(
                status_code=400,
                detail="Student ID and name are required."
            )
        if student.email:
            existing_email = (
                db.query(models.Student)
                .filter(func.lower(models.Student.email) == student.email.strip().lower())
                .first()
            )
            if existing_email:
                raise HTTPException(
                    status_code=409,
                    detail="This email address is already linked to a student."
                )

        new_student = models.Student(
            student_id=student.student_id.strip(),
            name=student.name.strip(),
            email=student.email.strip() if student.email else None,
            department=student.department,
            year=student.year,
            section_id=student.section_id
        )

        db.add(new_student)
        db.commit()
        db.refresh(new_student)

        return {
            "success": True,
            "message": "Student created successfully.",
            "student": {
                "id": new_student.id,
                "student_id": new_student.student_id,
                "name": new_student.name,
                "email": new_student.email,
                "department": new_student.department,
                "year": new_student.year,
                "section_id": new_student.section_id
            }
        }

    finally:
        db.close()


@app.get("/api/students")
def get_students():

    db = SessionLocal()

    try:

        students = (
            db.query(models.Student)
            .order_by(models.Student.id)
            .all()
        )

        return {
            "success": True,
            "count": len(students),
            "students": [
                {
                    "id": student.id,
                    "student_id": student.student_id,
                    "name": student.name,
                    "email": student.email,
                    "department": student.department,
                    "year": student.year,
                    "section_id": student.section_id,
                    "section_name": (
                        student.section.name
                        if student.section
                        else None
                    )
                }
                for student in students
            ]
        }

    finally:
        db.close()


@app.put("/api/students/{student_record_id}")
def update_student(student_record_id: int, student: StudentCreate):
    db = SessionLocal()
    try:
        existing_student = (
            db.query(models.Student)
            .filter(models.Student.id == student_record_id)
            .first()
        )
        if not existing_student:
            raise HTTPException(status_code=404, detail="Student not found.")
        section = (
            db.query(models.Section)
            .filter(models.Section.id == student.section_id)
            .first()
        )
        if not section:
            raise HTTPException(status_code=404, detail="Section not found.")
        student_code = student.student_id.strip()
        student_name = student.name.strip()
        if not student_code or not student_name:
            raise HTTPException(
                status_code=400,
                detail="Student ID and name are required."
            )
        duplicate_id = (
            db.query(models.Student)
            .filter(
                models.Student.student_id == student_code,
                models.Student.id != student_record_id
            )
            .first()
        )
        if duplicate_id:
            raise HTTPException(status_code=409, detail="Student ID already exists.")
        email = student.email.strip() if student.email else None
        if email:
            duplicate_email = (
                db.query(models.Student)
                .filter(
                    func.lower(models.Student.email) == email.lower(),
                    models.Student.id != student_record_id
                )
                .first()
            )
            if duplicate_email:
                raise HTTPException(
                    status_code=409,
                    detail="This email address is already linked to another student."
                )
        existing_student.student_id = student_code
        existing_student.name = student_name
        existing_student.email = email
        existing_student.department = student.department.strip() if student.department else None
        existing_student.year = student.year.strip() if student.year else None
        existing_student.section_id = student.section_id
        db.commit()
        db.refresh(existing_student)
        return {
            "success": True,
            "student": {
                "id": existing_student.id,
                "student_id": existing_student.student_id,
                "name": existing_student.name,
                "email": existing_student.email,
                "department": existing_student.department,
                "year": existing_student.year,
                "section_id": existing_student.section_id,
                "section_name": section.name
            }
        }
    finally:
        db.close()


@app.delete("/api/students/{student_record_id}")
def delete_student(student_record_id: int):
    db = SessionLocal()
    try:
        student = (
            db.query(models.Student)
            .filter(models.Student.id == student_record_id)
            .first()
        )
        if not student:
            raise HTTPException(status_code=404, detail="Student not found.")

        db.query(models.Submission).filter(
            models.Submission.student_record_id == student_record_id
        ).update(
            {models.Submission.student_record_id: None},
            synchronize_session=False
        )
        db.delete(student)
        db.commit()
        return {
            "success": True,
            "message": "Student account deleted. Existing submissions were preserved."
        }
    finally:
        db.close()


@app.get("/api/students/{student_id}")
def get_student(student_id: int):

    db = SessionLocal()

    try:

        student = (
            db.query(models.Student)
            .filter(
                models.Student.id == student_id
            )
            .first()
        )

        if not student:
            raise HTTPException(
                status_code=404,
                detail="Student not found."
            )

        return {
            "success": True,
            "student": {
                "id": student.id,
                "student_id": student.student_id,
                "name": student.name,
                "email": student.email,
                "department": student.department,
                "year": student.year,
                "section_id": student.section_id,
                "section_name": (
                    student.section.name
                    if student.section
                    else None
                )
            }
        }

    finally:
        db.close()


@app.get("/api/portal/students/resolve")
def resolve_portal_student(email: str):
    db = SessionLocal()
    try:
        students = (
            db.query(models.Student)
            .filter(func.lower(func.trim(models.Student.email)) == email.strip().lower())
            .all()
        )
        if not students:
            raise HTTPException(
                status_code=404,
                detail="No student account is linked to this email. Contact your administrator."
            )
        if len(students) > 1:
            raise HTTPException(
                status_code=409,
                detail="This email is linked to multiple student accounts. Contact your administrator."
            )
        student = students[0]
        return {
            "success": True,
            "student": {
                "student_id": student.student_id,
                "name": student.name,
                "email": student.email,
                "department": student.department,
                "year": student.year,
                "section_id": student.section_id,
                "section_name": student.section.name
            }
        }
    finally:
        db.close()


@app.get("/api/portal/students/{student_code}")
def get_portal_student_profile(student_code: str):
    db = SessionLocal()
    try:
        student = (
            db.query(models.Student)
            .filter(models.Student.student_id == student_code)
            .first()
        )
        if not student:
            raise HTTPException(status_code=404, detail="Student account not found.")
        return {
            "success": True,
            "student": {
                "student_id": student.student_id,
                "name": student.name,
                "email": student.email,
                "department": student.department,
                "year": student.year,
                "section_id": student.section_id,
                "section_name": student.section.name
            }
        }
    finally:
        db.close()


@app.put("/api/portal/students/{student_code}")
def update_portal_student_profile(
    student_code: str,
    profile: StudentProfileUpdate
):
    db = SessionLocal()
    try:
        student = (
            db.query(models.Student)
            .filter(models.Student.student_id == student_code)
            .first()
        )
        if not student:
            raise HTTPException(status_code=404, detail="Student account not found.")
        if not profile.name.strip():
            raise HTTPException(status_code=400, detail="Name is required.")
        if profile.email and "@" not in profile.email:
            raise HTTPException(status_code=400, detail="Enter a valid email address.")
        if profile.email:
            duplicate = (
                db.query(models.Student)
                .filter(
                    models.Student.email.ilike(profile.email.strip()),
                    models.Student.id != student.id
                )
                .first()
            )
            if duplicate:
                raise HTTPException(
                    status_code=409,
                    detail="This email address is already linked to another student."
                )
        student.name = profile.name.strip()
        student.email = profile.email.strip() if profile.email else None
        student.department = profile.department.strip() if profile.department else None
        student.year = profile.year.strip() if profile.year else None
        db.commit()
        db.refresh(student)
        return {
            "success": True,
            "student": {
                "student_id": student.student_id,
                "name": student.name,
                "email": student.email,
                "department": student.department,
                "year": student.year,
                "section_id": student.section_id,
                "section_name": student.section.name
            }
        }
    finally:
        db.close()
        # =========================================================
# EXCEL ROSTER IMPORT
# =========================================================

@app.post("/api/roster/import")
async def import_roster(
    file: UploadFile = File(...),
    faculty_name: str | None = Form(None)
):

    # -----------------------------------------------------
    # Validate file
    # -----------------------------------------------------

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file was selected."
        )

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Please upload an Excel .xlsx file."
        )

    # -----------------------------------------------------
    # Read Excel workbook
    # -----------------------------------------------------

    try:

        contents = await file.read()

        from io import BytesIO

        workbook = openpyxl.load_workbook(
            BytesIO(contents),
            data_only=True
        )

        worksheet = workbook.active

    except Exception as e:

        raise HTTPException(
            status_code=400,
            detail=f"Unable to read Excel file: {str(e)}"
        )

    # -----------------------------------------------------
    # Read header row
    # -----------------------------------------------------

    headers = []

    for cell in worksheet[1]:

        if cell.value is None:
            headers.append("")
        else:
            headers.append(
                str(cell.value).strip()
            )

    # Remove empty headers from the end
    while headers and headers[-1] == "":
        headers.pop()

    required_columns = [
        "Student ID",
        "Student Name",
        "Section"
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in headers
    ]

    if missing_columns:

        raise HTTPException(
            status_code=400,
            detail={
                "message": "Required columns are missing.",
                "missing_columns": missing_columns,
                "required_columns": required_columns
            }
        )

    # -----------------------------------------------------
    # Optional columns
    # -----------------------------------------------------

    optional_columns = [
        "Email",
        "Department",
        "Year"
    ]

    # -----------------------------------------------------
    # Convert Excel rows into dictionaries
    # -----------------------------------------------------

    rows = []

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=2,
            values_only=True
        ),
        start=2
    ):

        # Ignore completely empty rows
        if all(
            value is None or str(value).strip() == ""
            for value in row
        ):
            continue

        row_data = {}

        for index, header in enumerate(headers):

            if index < len(row):
                value = row[index]
            else:
                value = None

            if value is not None:
                value = str(value).strip()

            row_data[header] = value

        rows.append(
            {
                "row_number": row_number,
                "data": row_data
            }
        )

    if not rows:

        raise HTTPException(
            status_code=400,
            detail="The Excel file contains no student records."
        )

    # -----------------------------------------------------
    # Validate all rows BEFORE touching database
    # -----------------------------------------------------

    validation_errors = []

    student_ids_in_file = set()

    for item in rows:

        row_number = item["row_number"]
        data = item["data"]

        student_id = data.get("Student ID")
        student_name = data.get("Student Name")
        section_name = data.get("Section")

        # Required values
        if not student_id:
            validation_errors.append(
                f"Row {row_number}: Student ID is missing."
            )

        if not student_name:
            validation_errors.append(
                f"Row {row_number}: Student Name is missing."
            )

        if not section_name:
            validation_errors.append(
                f"Row {row_number}: Section is missing."
            )

        # Duplicate inside Excel
        if student_id:

            if student_id in student_ids_in_file:

                validation_errors.append(
                    f"Row {row_number}: Duplicate Student ID "
                    f"'{student_id}' in Excel file."
                )

            student_ids_in_file.add(student_id)

    # Stop before database changes
    if validation_errors:

        raise HTTPException(
            status_code=400,
            detail={
                "message": "Roster validation failed.",
                "errors": validation_errors
            }
        )

    # -----------------------------------------------------
    # Database operation
    # -----------------------------------------------------

    db = SessionLocal()

    try:

        # -------------------------------------------------
        # Check duplicate IDs against existing database
        # -------------------------------------------------

        existing_students = (
            db.query(models.Student)
            .filter(
                models.Student.student_id.in_(
                    list(student_ids_in_file)
                )
            )
            .all()
        )

        if existing_students:

            existing_ids = [
                student.student_id
                for student in existing_students
            ]

            raise HTTPException(
                status_code=400,
                detail={
                    "message":
                        "Some Student IDs already exist "
                        "in the database.",
                    "existing_student_ids":
                        existing_ids
                }
            )

        # -------------------------------------------------
        # Prepare sections
        # -------------------------------------------------

        section_names = sorted(
            {
                item["data"]["Section"]
                for item in rows
            }
        )

        section_objects = {}

        for section_name in section_names:

            section = (
                db.query(models.Section)
                .filter(
                    models.Section.name == section_name
                )
                .first()
            )

            if section:

                # If a faculty was supplied and the
                # section already belongs to another
                # faculty, reject the import.

                if (
                    faculty_name
                    and section.faculty_name
                    and section.faculty_name != faculty_name
                ):

                    raise HTTPException(
                        status_code=400,
                        detail=(
                            f"Section '{section_name}' "
                            f"is already assigned to "
                            f"{section.faculty_name}."
                        )
                    )

                # Add faculty if section doesn't have one
                if (
                    faculty_name
                    and not section.faculty_name
                ):

                    section.faculty_name = faculty_name

            else:

                section = models.Section(
                    name=section_name,
                    faculty_name=faculty_name
                )

                db.add(section)
                db.flush()

            section_objects[section_name] = section

        # -------------------------------------------------
        # Create students
        # -------------------------------------------------

        imported_students = []

        for item in rows:

            data = item["data"]

            student = models.Student(
                student_id=data["Student ID"],
                name=data["Student Name"],
                email=data.get("Email"),
                department=data.get("Department"),
                year=data.get("Year"),
                section_id=section_objects[
                    data["Section"]
                ].id
            )

            db.add(student)

            imported_students.append(
                {
                    "student_id":
                        student.student_id,
                    "name":
                        student.name,
                    "section":
                        data["Section"]
                }
            )

        # -------------------------------------------------
        # Commit everything together
        # -------------------------------------------------

        db.commit()

        return {
            "success": True,
            "message":
                "Roster imported successfully.",
            "filename":
                file.filename,
            "sections_created_or_updated":
                len(section_objects),
            "students_imported":
                len(imported_students),
            "students":
                imported_students
        }

    except HTTPException:
        db.rollback()
        raise

    except Exception as e:

        db.rollback()

        raise HTTPException(
            status_code=500,
            detail=(
                "Roster import failed. "
                f"No changes were saved. Error: {str(e)}"
            )
        )

    finally:

        db.close()
        # =========================================================
# CLASS ROSTER STATUS API
# =========================================================

@app.get("/api/sections/{section_id}/roster")
def get_section_roster(
    section_id: int,
    assignment_id: int | None = None
):

    db = SessionLocal()

    try:

        # -------------------------------------------------
        # Find section
        # -------------------------------------------------

        section = (
            db.query(models.Section)
            .filter(
                models.Section.id == section_id
            )
            .first()
        )

        if not section:

            raise HTTPException(
                status_code=404,
                detail="Section not found."
            )

        assignment = None
        if assignment_id is not None:
            assignment = (
                db.query(models.Assignment)
                .filter(
                    models.Assignment.id == assignment_id,
                    models.Assignment.section_id == section_id
                )
                .first()
            )
            if not assignment:
                raise HTTPException(
                    status_code=404,
                    detail="Assignment not found in this section."
                )

        # -------------------------------------------------
        # Get all students in this section
        # -------------------------------------------------

        students = (
            db.query(models.Student)
            .filter(
                models.Student.section_id == section_id
            )
            .order_by(models.Student.name)
            .all()
        )

        submission_query = db.query(models.Submission).filter(
            models.Submission.student_record_id.in_(
                [student.id for student in students]
            )
        )
        if assignment:
            submission_query = submission_query.filter(
                models.Submission.assignment_id == assignment.id
            )
        submissions = submission_query.order_by(
            models.Submission.submitted_at.desc()
        ).all()
        latest_submissions = {}
        for item in submissions:
            latest_submissions.setdefault(item.student_record_id, item)

        submission_ids = [item.id for item in latest_submissions.values()]
        evaluations = (
            db.query(models.Evaluation)
            .filter(models.Evaluation.submission_id.in_(submission_ids))
            .order_by(models.Evaluation.evaluated_at.desc())
            .all()
        )
        latest_evaluations = {}
        for item in evaluations:
            latest_evaluations.setdefault(item.submission_id, item)

        roster = []

        submitted_count = 0
        evaluated_count = 0
        pending_count = 0
        not_submitted_count = 0

        # -------------------------------------------------
        # Check every student
        # -------------------------------------------------

        for student in students:

            # Find latest submission for this student
            submission = latest_submissions.get(student.id)

            # -------------------------------------------------
            # Student has NOT submitted
            # -------------------------------------------------

            if not submission:

                submission_status = "NOT SUBMITTED"
                evaluation_status = None

                not_submitted_count += 1

                roster.append(
                    {
                        "student_id":
                            student.student_id,

                        "student_name":
                            student.name,

                        "email":
                            student.email,

                        "department":
                            student.department,

                        "year":
                            student.year,

                        "submission_status":
                            submission_status,

                        "evaluation_status":
                            evaluation_status,

                        "submission_id":
                            None,

                        "assignment_title":
                            assignment.title if assignment else None,

                        "subject":
                            None,

                        "submitted_at":
                            None,

                        "marks":
                            None,

                        "similarity_percentage":
                            None
                    }
                )

                continue

            # -------------------------------------------------
            # Student HAS submitted
            # -------------------------------------------------

            submitted_count += 1

            # Find latest evaluation
            evaluation = latest_evaluations.get(submission.id)

            # -------------------------------------------------
            # Submission evaluated
            # -------------------------------------------------

            if evaluation:

                evaluation_status = "EVALUATED"

                evaluated_count += 1

            # -------------------------------------------------
            # Submission waiting for evaluation
            # -------------------------------------------------

            else:

                evaluation_status = "PENDING"

                pending_count += 1

            roster.append(
                {
                    "student_id":
                        student.student_id,

                    "student_name":
                        student.name,

                    "email":
                        student.email,

                    "department":
                        student.department,

                    "year":
                        student.year,

                    "submission_status":
                        "SUBMITTED",

                    "evaluation_status":
                        evaluation_status,

                    "submission_id":
                        submission.id,

                    "assignment_title":
                        submission.assignment_title,

                    "subject":
                        submission.subject,

                    "submitted_at":
                        submission.submitted_at,

                    "marks":
                        evaluation.marks
                        if evaluation
                        else None,

                    "similarity_percentage":
                        evaluation.similarity_percentage
                        if evaluation
                        else None
                }
            )

        # -------------------------------------------------
        # Return complete roster
        # -------------------------------------------------

        return {
            "success": True,

            "section": {
                "id":
                    section.id,

                "name":
                    section.name,

                "faculty_name":
                    section.faculty_name
            },

            "summary": {
                "total_students":
                    len(students),

                "submitted":
                    submitted_count,

                "not_submitted":
                    not_submitted_count,

                "evaluated":
                    evaluated_count,

                "pending_evaluation":
                    pending_count
            },

            "students":
                roster
        }

    finally:

        db.close()


# =========================================================
# EXCEL ROSTER EXPORT
# =========================================================

@app.get("/api/sections/{section_id}/roster/export")
def export_section_roster(
    section_id: int,
    assignment_id: int | None = None
):

    db = SessionLocal()

    try:
        section = (
            db.query(models.Section)
            .filter(models.Section.id == section_id)
            .first()
        )

        if not section:
            raise HTTPException(
                status_code=404,
                detail="Section not found."
            )

        selected_assignment = None
        if assignment_id is not None:
            selected_assignment = (
                db.query(models.Assignment)
                .filter(
                    models.Assignment.id == assignment_id,
                    models.Assignment.section_id == section_id
                )
                .first()
            )
            if not selected_assignment:
                raise HTTPException(
                    status_code=404,
                    detail="Assignment not found in this section."
                )

        students = (
            db.query(models.Student)
            .filter(models.Student.section_id == section_id)
            .order_by(models.Student.student_id.asc())
            .all()
        )

        submission_query = db.query(models.Submission).filter(
            models.Submission.student_record_id.in_(
                [student.id for student in students]
            )
        )
        if selected_assignment:
            submission_query = submission_query.filter(
                models.Submission.assignment_id == selected_assignment.id
            )
        submissions = submission_query.order_by(
            models.Submission.submitted_at.desc()
        ).all()
        latest_submissions = {}
        for item in submissions:
            latest_submissions.setdefault(item.student_record_id, item)

        submission_ids = [item.id for item in latest_submissions.values()]
        evaluations = (
            db.query(models.Evaluation)
            .filter(models.Evaluation.submission_id.in_(submission_ids))
            .order_by(models.Evaluation.evaluated_at.desc())
            .all()
        )
        latest_evaluations = {}
        for item in evaluations:
            latest_evaluations.setdefault(item.submission_id, item)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Assignment Status" if selected_assignment else "Class Roster"

        headers = [
            "Student ID",
            "Student Name",
            "Section",
            "Assignment",
            "Submission Status",
            "Submission Date",
            "Evaluation Status",
            "Marks",
            "Similarity %",
            "Faculty Feedback"
        ]

        worksheet.append(headers)

        for cell in worksheet[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="172033"
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        for student in students:

            submission = latest_submissions.get(student.id)
            evaluation = (
                latest_evaluations.get(submission.id)
                if submission
                else None
            )

            if not submission:
                submission_status = "NOT SUBMITTED"
                evaluation_status = "NOT SUBMITTED"
                assignment_title = selected_assignment.title if selected_assignment else "—"
                submission_date = "—"
                marks = "—"
                similarity = "—"
                feedback = "—"

            else:
                submission_status = "SUBMITTED"
                assignment_title = submission.assignment_title or "—"
                submission_date = (
                    submission.submitted_at.strftime("%Y-%m-%d %H:%M")
                    if submission.submitted_at
                    else "—"
                )

                if evaluation:
                    evaluation_status = "EVALUATED"
                    marks = (
                        f"{evaluation.marks:.1f}"
                        if evaluation.marks is not None
                        else "—"
                    )
                    similarity = (
                        f"{evaluation.similarity_percentage:.1f}%"
                        if evaluation.similarity_percentage is not None
                        else "—"
                    )
                    feedback = evaluation.feedback or "—"

                else:
                    evaluation_status = "PENDING"
                    marks = "—"
                    similarity = "—"
                    feedback = "—"

            worksheet.append([
                student.student_id,
                student.name,
                section.name,
                assignment_title,
                submission_status,
                submission_date,
                evaluation_status,
                marks,
                similarity,
                feedback
            ])

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        column_widths = {
            "A": 16,
            "B": 24,
            "C": 16,
            "D": 30,
            "E": 20,
            "F": 22,
            "G": 20,
            "H": 12,
            "I": 16,
            "J": 45
        }

        for column, width in column_widths.items():
            worksheet.column_dimensions[column].width = width

        if worksheet.max_row >= 2:
            for row in worksheet.iter_rows(
                min_row=2,
                max_row=worksheet.max_row
            ):
                row[0].alignment = Alignment(horizontal="center")
                row[2].alignment = Alignment(horizontal="center")
                row[4].alignment = Alignment(horizontal="center")
                row[5].alignment = Alignment(horizontal="center")
                row[6].alignment = Alignment(horizontal="center")
                row[7].alignment = Alignment(horizontal="center")
                row[8].alignment = Alignment(horizontal="center")

        output = BytesIO()

        workbook.save(output)
        output.seek(0)

        safe_section_name = (
            section.name
            .replace(" ", "_")
            .replace("/", "_")
            .replace("\\", "_")
        )

        safe_assignment_name = (
            "_" + "".join(
                character if character.isalnum() else "_"
                for character in selected_assignment.title
            ).strip("_")[:80]
            if selected_assignment
            else "_Class_Roster"
        )
        filename = f"{safe_section_name}{safe_assignment_name}.xlsx"

        return Response(
            content=output.getvalue(),
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}"'
            }
        )

    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Excel export failed: {str(e)}"
        )

    finally:
        db.close()


# =========================================================
# SECTION ASSIGNMENTS
# =========================================================

@app.post("/api/assignments")
def create_assignment(request: AssignmentCreate):
    db = SessionLocal()
    try:
        section = (
            db.query(models.Section)
            .filter(models.Section.id == request.section_id)
            .first()
        )
        if not section:
            raise HTTPException(status_code=404, detail="Section not found.")
        if not request.title.strip() or not request.subject.strip():
            raise HTTPException(
                status_code=400,
                detail="Assignment title and subject are required."
            )

        assignment = models.Assignment(
            section_id=section.id,
            title=request.title.strip(),
            subject=request.subject.strip(),
            description=request.description,
            due_date=(
                request.due_date.astimezone(timezone.utc).replace(tzinfo=None)
                if request.due_date and request.due_date.tzinfo
                else request.due_date
            ),
            faculty_name=request.faculty_name or section.faculty_name
        )
        db.add(assignment)
        db.flush()
        section_student_ids = db.query(models.Student.id).filter(
            models.Student.section_id == section.id
        )
        db.query(models.Submission).filter(
            models.Submission.assignment_id.is_(None),
            models.Submission.student_record_id.in_(section_student_ids),
            func.lower(func.trim(models.Submission.assignment_title))
            == func.lower(func.trim(assignment.title)),
            func.lower(func.trim(models.Submission.subject))
            == func.lower(func.trim(assignment.subject))
        ).update(
            {models.Submission.assignment_id: assignment.id},
            synchronize_session=False
        )
        db.commit()
        db.refresh(assignment)
        return {
            "success": True,
            "assignment": {
                "id": assignment.id,
                "section_id": assignment.section_id,
                "section_name": section.name,
                "title": assignment.title,
                "subject": assignment.subject,
                "description": assignment.description,
                "due_date": serialize_due_date(assignment.due_date),
                "faculty_name": assignment.faculty_name,
                "created_at": assignment.created_at.isoformat()
            }
        }
    finally:
        db.close()


@app.get("/api/assignments")
def get_assignments(section_id: int | None = None):
    db = SessionLocal()
    try:
        query = db.query(models.Assignment)
        if section_id is not None:
            query = query.filter(models.Assignment.section_id == section_id)
        assignments = query.order_by(models.Assignment.created_at.desc()).all()
        return {
            "success": True,
            "assignments": [
                {
                    "id": item.id,
                    "section_id": item.section_id,
                    "section_name": item.section.name,
                    "title": item.title,
                    "subject": item.subject,
                    "description": item.description,
                    "due_date": serialize_due_date(item.due_date),
                    "faculty_name": item.faculty_name,
                    "created_at": item.created_at.isoformat()
                }
                for item in assignments
            ]
        }
    finally:
        db.close()


@app.get("/api/roster/students/{student_code}/assignments")
def get_student_assignments(student_code: str):
    db = SessionLocal()
    try:
        student = (
            db.query(models.Student)
            .filter(models.Student.student_id == student_code)
            .first()
        )
        if not student:
            raise HTTPException(status_code=404, detail="Student ID not found in the roster.")

        assignments = (
            db.query(models.Assignment)
            .filter(models.Assignment.section_id == student.section_id)
            .order_by(models.Assignment.created_at.desc())
            .all()
        )
        assignment_rows = []
        for assignment in assignments:
            submission = (
                db.query(models.Submission)
                .filter(
                    models.Submission.student_record_id == student.id,
                    models.Submission.assignment_id == assignment.id
                )
                .order_by(models.Submission.submitted_at.desc())
                .first()
            )
            assignment_rows.append({
                "id": assignment.id,
                "title": assignment.title,
                "subject": assignment.subject,
                "description": assignment.description,
                "due_date": serialize_due_date(assignment.due_date),
                "submission_id": submission.id if submission else None,
                "submission_status": submission.status if submission else "NOT SUBMITTED"
            })

        return {
            "success": True,
            "student": {
                "student_id": student.student_id,
                "name": student.name,
                "section_id": student.section_id,
                "section_name": student.section.name
            },
            "assignments": assignment_rows
        }
    finally:
        db.close()


@app.get("/api/roster/students/{student_code}/submissions")
def get_student_submissions(student_code: str):
    db = SessionLocal()
    try:
        student = (
            db.query(models.Student)
            .filter(models.Student.student_id == student_code)
            .first()
        )
        if not student:
            raise HTTPException(status_code=404, detail="Student ID not found in the roster.")

        submissions = (
            db.query(models.Submission)
            .filter(models.Submission.student_record_id == student.id)
            .order_by(models.Submission.submitted_at.desc())
            .all()
        )
        results = []
        for submission in submissions:
            evaluation = (
                db.query(models.Evaluation)
                .filter(models.Evaluation.submission_id == submission.id)
                .order_by(models.Evaluation.evaluated_at.desc())
                .first()
            )
            results.append({
                "id": submission.id,
                "assignment_id": submission.assignment_id,
                "assignment_title": submission.assignment_title,
                "subject": submission.subject,
                "filename": submission.filename,
                "status": submission.status,
                "submitted_at": submission.submitted_at.isoformat() if submission.submitted_at else None,
                "evaluation": {
                    "marks": evaluation.marks,
                    "feedback": evaluation.feedback,
                    "similarity_percentage": evaluation.similarity_percentage,
                    "evaluated_at": evaluation.evaluated_at.isoformat() if evaluation.evaluated_at else None
                } if evaluation else None
            })
        return {"success": True, "student": student.name, "submissions": results}
    finally:
        db.close()
